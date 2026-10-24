from pathlib import Path

from docx import Document as WordDocument
from langchain_core.documents import Document
from pypdf import PdfWriter

from app.config import settings
from app.rag.ingestion import (
    extract_docx_text,
    extract_html_text,
    extract_pdf_text,
    load_documents,
    split_documents,
)


def test_extract_html_text_removes_scripts(tmp_path: Path):
    file_path = tmp_path / "agreement.doc"
    file_path.write_text(
        """
        <html>
            <body>
                <h1>Convenção Coletiva</h1>
                <p>Seguro de vida obrigatório.</p>
                <script>alert("ignored")</script>
            </body>
        </html>
        """,
        encoding="utf-8",
    )

    content = extract_html_text(file_path)

    assert "Convenção Coletiva" in content
    assert "Seguro de vida obrigatório." in content
    assert "ignored" not in content


def test_load_documents_preserves_source(tmp_path: Path, monkeypatch):
    file_path = tmp_path / "agreement.doc"
    file_path.write_text(
        "<html><body><p>Assistência médica obrigatória.</p></body></html>",
        encoding="utf-8",
    )

    monkeypatch.setattr(settings, "documents_dir", tmp_path)

    documents = load_documents()

    assert len(documents) == 1
    assert documents[0].metadata["source"] == "agreement.doc"
    assert "Assistência médica" in documents[0].page_content


def test_load_documents_ignores_unsupported_files(
    tmp_path: Path,
    monkeypatch,
):
    (tmp_path / "data.csv").write_text(
        "name,value",
        encoding="utf-8",
    )

    monkeypatch.setattr(settings, "documents_dir", tmp_path)

    documents = load_documents()

    assert documents == []


def test_split_documents_preserves_metadata():
    document = Document(
        page_content="Seguro de vida obrigatório. " * 100,
        metadata={"source": "agreement.doc"},
    )

    chunks = split_documents([document])

    assert len(chunks) > 1
    assert all(
        chunk.metadata["source"] == "agreement.doc"
        for chunk in chunks
    )
    assert all(chunk.page_content.strip() for chunk in chunks)


def test_extract_docx_text_includes_paragraphs_and_tables(
    tmp_path: Path,
):
    file_path = tmp_path / "agreement.docx"

    document = WordDocument()
    document.add_paragraph("Convenção Coletiva")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Benefício"
    table.cell(0, 1).text = "Seguro de vida"

    document.add_paragraph("Assistência médica obrigatória.")
    document.save(file_path)

    content = extract_docx_text(file_path)

    assert "Convenção Coletiva" in content
    assert "Benefício | Seguro de vida" in content
    assert "Assistência médica obrigatória." in content

    assert content.index("Convenção Coletiva") < content.index(
        "Benefício | Seguro de vida"
    )
    assert content.index("Benefício | Seguro de vida") < content.index(
        "Assistência médica obrigatória."
    )


def test_extract_pdf_text_handles_empty_pages(tmp_path: Path):
    file_path = tmp_path / "agreement.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)

    with file_path.open("wb") as output:
        writer.write(output)

    assert extract_pdf_text(file_path) == ""


def test_load_documents_supports_multiple_formats(
    tmp_path: Path,
    monkeypatch,
):
    (tmp_path / "agreement.txt").write_text(
        "Adicional noturno previsto.",
        encoding="utf-8",
    )

    (tmp_path / "agreement.doc").write_text(
        "<html><body><p>Seguro de vida obrigatório.</p></body></html>",
        encoding="utf-8",
    )

    word_document = WordDocument()
    word_document.add_paragraph("Vale-alimentação previsto.")
    word_document.save(tmp_path / "agreement.docx")

    monkeypatch.setattr(settings, "documents_dir", tmp_path)

    documents = load_documents()

    assert len(documents) == 3

    sources = {
        document.metadata["source"]
        for document in documents
    }

    assert sources == {
        "agreement.txt",
        "agreement.doc",
        "agreement.docx",
    }


def test_load_documents_ignores_empty_pdf(
    tmp_path: Path,
    monkeypatch,
):
    file_path = tmp_path / "empty.pdf"

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)

    with file_path.open("wb") as output:
        writer.write(output)

    monkeypatch.setattr(settings, "documents_dir", tmp_path)

    assert load_documents() == []